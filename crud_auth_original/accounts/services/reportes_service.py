from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet


class ReportesService:

    # ---------- EXCEL ----------
    @staticmethod
    def usuarios_a_excel(usuarios):
        wb = Workbook()
        ws = wb.active
        ws.title = "Usuarios"

        headers = ["ID", "Usuario", "Nombre", "Apellido", "Email",
                   "DNI", "Grupo", "Activo", "Staff", "Fecha registro"]
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="34495E")
        for col, _ in enumerate(headers, 1):
            c = ws.cell(row=1, column=col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")

        for u in usuarios:
            ws.append([
                u["id"], u["username"], u["nombre"], u["apellido"],
                u["email"], u["dni"], u["grupo"],
                "Sí" if u["is_active"] else "No",
                "Sí" if u["is_staff"] else "No",
                u["fecha_registro"],
            ])

        # Ancho automático
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 3

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # ---------- PDF ----------
    @staticmethod
    def usuarios_a_pdf(usuarios, titulo="Informe de Usuarios", filtros=None):
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=landscape(A4),
            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        )
        styles = getSampleStyleSheet()
        elementos = []

        elementos.append(Paragraph(f"<b>{titulo}</b>", styles["Title"]))
        elementos.append(Spacer(1, 0.4 * cm))

        if filtros:
            texto_filtros = " | ".join(f"{k}: {v}" for k, v in filtros.items() if v)
            if texto_filtros:
                elementos.append(Paragraph(f"<i>Filtros: {texto_filtros}</i>",
                                            styles["Normal"]))
                elementos.append(Spacer(1, 0.4 * cm))

        data = [["ID", "Usuario", "Nombre", "Apellido", "Email",
                 "Grupo", "Activo", "Fecha"]]
        for u in usuarios:
            data.append([
                u["id"], u["username"], u["nombre"], u["apellido"],
                u["email"], u["grupo"],
                "Sí" if u["is_active"] else "No",
                u["fecha_registro"],
            ])

        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#34495E")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F2F2F2")]),
        ]))
        elementos.append(tabla)

        doc.build(elementos)
        buffer.seek(0)
        return buffer